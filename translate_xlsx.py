#!/usr/bin/env python3
import sys
import argparse
from typing import Any, Dict

try:
    from deep_translator import GoogleTranslator
except Exception as exc:  # noqa: BLE001
    print("Failed to import deep_translator. Please ensure it's installed.", file=sys.stderr)
    raise

try:
    from openpyxl import load_workbook
except Exception as exc:  # noqa: BLE001
    print("Failed to import openpyxl. Please ensure it's installed.", file=sys.stderr)
    raise


def translate_string(text: str, translator: GoogleTranslator, cache: Dict[str, str]) -> str:
    if not text:
        return text
    key = text.strip()
    if key in cache:
        return cache[key]
    try:
        translated = translator.translate(text)
        cache[key] = translated
        return translated
    except Exception:
        # On any translation failure, return original text
        cache[key] = text
        return text


def translate_workbook(input_path: str, output_path: str) -> None:
    translator = GoogleTranslator(source='auto', target='en')
    translation_cache: Dict[str, str] = {}

    workbook = load_workbook(filename=input_path)

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value: Any = cell.value
                # Skip formulas, numbers, booleans, None
                if value is None:
                    continue
                if isinstance(value, (int, float)):
                    continue
                if isinstance(value, bool):
                    continue
                # openpyxl sets data_type 'f' for formulas
                if cell.data_type == 'f':
                    continue
                if isinstance(value, str):
                    stripped = value.strip()
                    if not stripped:
                        continue
                    # Skip likely URLs
                    if stripped.lower().startswith(("http://", "https://")):
                        continue
                    cell.value = translate_string(value, translator, translation_cache)

    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Translate all string cells in an Excel workbook to English.")
    parser.add_argument("input", help="Path to input .xlsx file")
    parser.add_argument("output", help="Path to output .xlsx file")
    args = parser.parse_args()

    translate_workbook(args.input, args.output)
    print(f"Translated workbook saved to: {args.output}")


if __name__ == "__main__":
    main()